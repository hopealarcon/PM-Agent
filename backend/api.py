import os
import json
import re
import io
import anthropic
from fastapi import FastAPI, HTTPException, UploadFile, File
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from typing import Optional
from dotenv import load_dotenv

from db import repositories
from schemas.plan import Session, ScopeDecision, Decision, Plan
from prompts.clarification import SYSTEM_PROMPT as CLARIFY_SYSTEM, build_clarification_prompt
from prompts.scope_review import SYSTEM_PROMPT as SCOPE_SYSTEM, build_epics_prompt, build_features_prompt
from prompts.plan_generation import SYSTEM_PROMPT as PLAN_SYSTEM, build_plan_prompt
from prompts.timeline import build_timeline_prompt
from prompts.release_analysis import build_release_analysis_prompt
from prompts.release_schedule import build_release_schedule_prompt

load_dotenv()

app = FastAPI()

app.add_middleware(
    CORSMiddleware,
    allow_origins=["*"],
    allow_credentials=False,
    allow_methods=["*"],
    allow_headers=["*"],
)

claude = anthropic.Anthropic(api_key="".join(os.environ["ANTHROPIC_API_KEY"].split()))


# --- Request/Response models ---

class StartProjectRequest(BaseModel):
    name: str
    brief: str

class ClarifyRequest(BaseModel):
    brief: str
    history: list[dict]

class ProposeEpicsRequest(BaseModel):
    brief: str
    clarifications: list[dict]
    document_context: Optional[str] = None

class ProposeFeaturesRequest(BaseModel):
    epic_title: str
    brief: str
    clarifications: list[dict]
    document_context: Optional[str] = None

class GeneratePlanRequest(BaseModel):
    name: str
    brief: str
    clarifications: list[dict]
    accepted_scope: dict
    scope_decisions: list[dict]
    document_context: Optional[str] = None

class TimelineRequest(BaseModel):
    resources: list[dict]
    start_date: str

class UpdateRecordRequest(BaseModel):
    table: str
    id: str
    fields: dict


class AnalyzeReleasesRequest(BaseModel):
    releases: list[dict]  # [{name, description, raw_tasks}]


class SaveReleasePlanRequest(BaseModel):
    name: str
    releases: list[dict]  # analyzed releases with tasks
    resources: list[dict]
    start_date: str


# --- Endpoints ---

@app.get("/api/projects")
def list_projects():
    return repositories.get_projects()


@app.get("/api/projects/{project_id}")
def get_project(project_id: str):
    plan = repositories.get_plan(project_id)
    if not plan:
        raise HTTPException(status_code=404, detail="Plan no encontrado")
    return plan


@app.post("/api/clarify")
def clarify(req: ClarifyRequest):
    round_num = len(req.history) + 1
    prompt = build_clarification_prompt(req.brief, req.history, round_num)

    response = claude.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=1024,
        system=CLARIFY_SYSTEM,
        messages=[{"role": "user", "content": prompt}]
    )
    text = response.content[0].text.strip()

    if "LISTO_PARA_PLANEAR" in text:
        return {"ready": True, "questions": None}

    return {"ready": False, "questions": text}


@app.post("/api/scope/epics")
def propose_epics(req: ProposeEpicsRequest):
    prompt = build_epics_prompt(req.brief, req.clarifications, req.document_context or "")
    response = claude.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=1024,
        system=SCOPE_SYSTEM,
        messages=[{"role": "user", "content": prompt}]
    )
    text = response.content[0].text.strip()

    epics = []
    for line in text.split("\n"):
        clean = re.sub(r'\*+', '', line.strip())  # remove markdown bold
        match = re.match(r'\[([A-Z])\]\s+(.+?)(?:\s+-\s+(.+))?$', clean)
        if match:
            epics.append({
                "key": match.group(1),
                "title": match.group(2).strip(),
                "description": match.group(3).strip() if match.group(3) else ""
            })

    return {"epics": epics}


@app.post("/api/scope/features")
def propose_features(req: ProposeFeaturesRequest):
    prompt = build_features_prompt(req.epic_title, req.brief, req.clarifications, req.document_context or "")
    response = claude.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=1024,
        system=SCOPE_SYSTEM,
        messages=[{"role": "user", "content": prompt}]
    )
    text = response.content[0].text.strip()

    features = []
    for line in text.split("\n"):
        clean = re.sub(r'\*+', '', line.strip())
        match = re.match(r'\[(\d+)\]\s+(.+?)(?:\s+-\s+(.+?))?(?:\s+\((must-have|should-have|could-have)\))?$', clean)
        if match:
            features.append({
                "title": match.group(2).strip(),
                "description": match.group(3).strip() if match.group(3) else "",
                "priority": match.group(4) or "should-have"
            })

    return {"features": features}


@app.post("/api/projects")
def generate_and_save(req: GeneratePlanRequest):
    try:
        prompt = build_plan_prompt(req.brief, req.clarifications, req.accepted_scope, req.document_context or "")
        response = claude.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=16000,
            system=PLAN_SYSTEM,
            messages=[{"role": "user", "content": prompt}]
        )
        text = response.content[0].text.strip()
        text = re.sub(r'^```[a-z]*\n?', '', text)
        text = re.sub(r'\n?```$', '', text)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Claude error: {e}")

    try:
        plan = Plan(**json.loads(text))
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"JSON parse error: {e} | raw: {text[:300]}")

    try:
        session = Session(
            project_name=req.name,
            brief_raw=req.brief,
            clarification_history=req.clarifications,
            scope_decisions=[ScopeDecision(**d) for d in req.scope_decisions],
            plan=plan
        )
        project_id = repositories.save_project(req.name)
        repositories.save_session(project_id, session)
        repositories.save_scope_decisions(project_id, session.scope_decisions)
        repositories.save_plan(project_id, plan)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"DB error: {e}")

    return {"project_id": project_id}


@app.patch("/api/update")
def update_record(req: UpdateRecordRequest):
    allowed = {"epics", "features", "tasks", "milestones", "risks"}
    if req.table not in allowed:
        raise HTTPException(status_code=400, detail="Tabla no permitida")
    try:
        from db.client import get_client
        client = get_client()
        client.update(req.table, req.id, req.fields)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error actualizando: {e}")
    return {"ok": True}


@app.post("/api/timeline/{project_id}")
def generate_timeline(project_id: str, req: TimelineRequest):
    plan = repositories.get_plan(project_id)
    if not plan:
        raise HTTPException(status_code=404, detail="Plan no encontrado")

    prompt = build_timeline_prompt(plan, req.resources, req.start_date)
    try:
        response = claude.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8000,
            system="Eres un project manager experto en planificacion de recursos y timelines. Asignas trabajo considerando dependencias, capacidad del equipo y prioridades. Respondes siempre en JSON valido.",
            messages=[{"role": "user", "content": prompt}]
        )
        text = response.content[0].text.strip()
        text = re.sub(r'^```[a-z]*\n?', '', text)
        text = re.sub(r'\n?```$', '', text)
        data = json.loads(text)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando timeline: {e}")

    try:
        repositories.save_timeline(project_id, req.resources, req.start_date, data["entries"])
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"DB error: {e}")

    return {"entries": data["entries"]}


@app.get("/api/timeline/{project_id}")
def get_timeline(project_id: str):
    timeline = repositories.get_timeline(project_id)
    if not timeline:
        return {"entries": [], "resources": [], "start_date": None}
    return timeline


@app.post("/api/extract-tasks")
async def extract_tasks(file: UploadFile = File(...)):
    content = await file.read()
    filename = file.filename or ""

    try:
        if filename.lower().endswith(".xlsx"):
            import openpyxl
            wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
            lines = []
            for sheet in wb.worksheets:
                for row in sheet.iter_rows():
                    cells = [str(cell.value) if cell.value is not None else "" for cell in row]
                    line = "\t".join(cells).rstrip("\t")
                    if line.strip():
                        lines.append(line)
            text = "\n".join(lines)
        elif filename.lower().endswith(".csv"):
            text = content.decode("utf-8", errors="ignore")
        else:
            text = content.decode("utf-8", errors="ignore")
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"No se pudo leer el archivo: {e}")

    if not text.strip():
        raise HTTPException(status_code=422, detail="El archivo no contiene texto extraible")

    return {"text": text[:8000]}


@app.get("/api/releases")
def list_release_plans():
    return repositories.get_release_plans()


@app.post("/api/releases/analyze")
def analyze_releases(req: AnalyzeReleasesRequest):
    prompt = build_release_analysis_prompt(req.releases)
    try:
        response = claude.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8000,
            system="Responde SIEMPRE con JSON valido y completo. Sin comentarios, sin texto adicional, sin trailing commas. Cierra todos los brackets y llaves correctamente.",
            messages=[{"role": "user", "content": prompt}]
        )
        text = response.content[0].text.strip()
        text = re.sub(r'^```[a-z]*\n?', '', text)
        text = re.sub(r'\n?```$', '', text)
        data = json.loads(text)
    except json.JSONDecodeError as e:
        raise HTTPException(status_code=500, detail=f"Error analizando releases: JSON invalido - {e}")
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error analizando releases: {e}")

    return data


@app.post("/api/releases")
def save_release_plan(req: SaveReleasePlanRequest):
    prompt = build_release_schedule_prompt(req.releases, req.resources, req.start_date)
    try:
        response = claude.messages.create(
            model="claude-sonnet-4-6",
            max_tokens=8000,
            messages=[{"role": "user", "content": prompt}]
        )
        text = response.content[0].text.strip()
        text = re.sub(r'^```[a-z]*\n?', '', text)
        text = re.sub(r'\n?```$', '', text)
        data = json.loads(text)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"Error generando schedule: {e}")

    try:
        plan_id = repositories.save_release_plan(
            name=req.name,
            releases=req.releases,
            resources=req.resources,
            start_date=req.start_date,
            entries=data["entries"],
            go_live_dates=data["go_live_dates"],
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"DB error: {e}")

    return {"plan_id": plan_id}


@app.get("/api/releases/{plan_id}")
def get_release_plan(plan_id: str):
    plan = repositories.get_release_plan(plan_id)
    if not plan:
        raise HTTPException(status_code=404, detail="Plan no encontrado")
    return plan


@app.post("/api/extract-document")
async def extract_document(file: UploadFile = File(...)):
    content = await file.read()
    filename = file.filename or ""

    try:
        if filename.lower().endswith(".pdf"):
            import pypdf
            reader = pypdf.PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        elif filename.lower().endswith(".docx"):
            import docx
            doc = docx.Document(io.BytesIO(content))
            text = "\n".join(p.text for p in doc.paragraphs if p.text.strip())
        else:
            text = content.decode("utf-8", errors="ignore")
    except Exception as e:
        raise HTTPException(status_code=422, detail=f"No se pudo leer el archivo: {e}")

    if not text.strip():
        raise HTTPException(status_code=422, detail="El archivo no contiene texto extraible")

    text = text[:12000]

    response = claude.messages.create(
        model="claude-sonnet-4-6",
        max_tokens=1024,
        system="Eres un asistente de product management. Tu tarea es leer un documento de requerimientos y generar un resumen ejecutivo en español que sirva como brief para un proyecto de software. El brief debe incluir: objetivo del proyecto, usuarios objetivo, principales funcionalidades, y cualquier restriccion o contexto relevante. Responde SOLO con el brief, sin titulos ni formato markdown.",
        messages=[{"role": "user", "content": f"Documento:\n\n{text}"}]
    )
    brief = response.content[0].text.strip()
    return {"brief": brief, "document_text": text}


@app.get("/health")
def health():
    return {"status": "ok"}
